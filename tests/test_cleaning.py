"""
Tests for the raw -> cleaned staging pipeline (etl/run_cleaning_pipeline.py).

Unit tests use small in-memory fixtures. Integration tests run the real
pipeline on the real raw files into a temp directory (see conftest.py) and
then run every validation check in etl/validate_cleaned_data.py against it.

    .venv\\Scripts\\python.exe -m pytest tests -v
"""

import math

import openpyxl
import pandas as pd
import pytest

import cleaning_common as cc
import validate_cleaned_data as vcd
from cleaning_common import OUTPUT_FILES


# ---------------------------------------------------------------------------
# Unit tests: helpers and marker handling
# ---------------------------------------------------------------------------

class TestTextAndNumbers:
    @pytest.mark.parametrize("raw,expected", [
        ("Household total income groups (14)", "household_total_income_groups"),
        ("Statistics (3C)", "statistics"),
        ("C2_COUNT_MEN+", "c2_count_men_plus"),
        ("Shelter-cost-to-income ratio (5)", "shelter_cost_to_income_ratio"),
    ])
    def test_snake_case(self, raw, expected):
        assert cc.snake_case(raw) == expected

    def test_normalize_text_trims_and_collapses(self):
        assert cc.normalize_text("Total - Dwelling condition ") == "Total - Dwelling condition"
        assert cc.normalize_text("Oct-23\nto\nOct-24") == "Oct-23 to Oct-24"
        assert cc.normalize_text("   ") is None
        assert cc.normalize_text(None) is None

    def test_parse_number_strips_thousands_separator(self):
        assert cc.parse_number("1,058") == 1058.0
        assert cc.parse_number(4.3) == 4.3
        with pytest.raises(ValueError):
            cc.parse_number("abc")


class TestMarkers:
    def test_cmhc_rms_suppressed(self):
        value, status = cc.parse_cmhc_rms_value("**")
        assert math.isnan(value) and status == "suppressed"

    def test_cmhc_rms_change_not_significant(self):
        value, status = cc.parse_cmhc_rms_value("++")
        assert math.isnan(value) and status == "change_not_significant"

    def test_cmhc_rms_published(self):
        assert cc.parse_cmhc_rms_value("1,325") == (1325.0, "published")

    def test_cmhc_rms_unknown_marker_raises(self):
        with pytest.raises(ValueError):
            cc.parse_cmhc_rms_value("??")

    def test_cmhc_starts_dash_is_zero_with_flag(self):
        assert cc.parse_cmhc_starts_value("-") == (0.0, "nil_or_zero")

    def test_cmhc_starts_not_available(self):
        value, flag = cc.parse_cmhc_starts_value("**")
        assert math.isnan(value) and flag == "not_available"

    @pytest.mark.parametrize("raw,expected", [("↑", "significant_increase"), ("↓", "significant_decrease"),
                                              ("-", "not_significant"), (None, None)])
    def test_cmhc_trend(self, raw, expected):
        assert cc.parse_cmhc_trend(raw) == expected

    def test_cmhc_reliability(self):
        assert cc.parse_cmhc_reliability("b") == "b"
        assert cc.parse_cmhc_reliability(" ") is None
        with pytest.raises(ValueError):
            cc.parse_cmhc_reliability("z")

    @pytest.mark.parametrize("raw,expected", [("", "published"), ("x", "suppressed"), ("...", "not_applicable"),
                                              (" r", "revised"), ("F", "too_unreliable")])
    def test_statcan_status(self, raw, expected):
        assert cc.statcan_status(raw) == expected

    def test_statcan_unknown_symbol_raises(self):
        with pytest.raises(ValueError):
            cc.statcan_status("Q")


class TestLayoutGuard:
    def test_assert_cells_passes_on_expected_layout(self):
        ws = openpyxl.Workbook().active
        ws["A7"], ws["B7"] = "Zone", "Oct-24"
        cc.assert_cells(ws, {"A7": "Zone", "B7": "Oct-24"})

    def test_assert_cells_raises_on_different_layout(self):
        ws = openpyxl.Workbook().active
        ws["A7"], ws["B7"] = "Zone", "Oct-25"
        with pytest.raises(ValueError, match="Unexpected layout"):
            cc.assert_cells(ws, {"A7": "Zone", "B7": "Oct-24"})


class TestShelterCleaningRules:
    def _raw(self, value, status):
        import clean_shelter_cost as m
        row = {c: "" for c in ["REF_DATE", "GEO", "DGUID", *m.DIMENSIONS, *m.CONSTANT_COLUMNS,
                                "COORDINATE", "VALUE", "STATUS"]}
        row.update({"REF_DATE": "2021", "GEO": "Windsor (CMA), Ont.", "DGUID": "2021S0503559",
                    "COORDINATE": "96.13.1.1.1.1.5.2", "VALUE": value, "STATUS": status})
        for col in m.DIMENSIONS:
            row[col] = "Some category "
        return pd.DataFrame([row])

    def test_not_applicable_zero_becomes_nan(self):
        import clean_shelter_cost as m
        out = m.clean_rows(self._raw("0", "..."))
        assert math.isnan(out.loc[0, "value"]) and out.loc[0, "value_status"] == "not_applicable"
        assert out.loc[0, "raw_value"] == "0"

    def test_suppressed_blank_becomes_nan(self):
        import clean_shelter_cost as m
        out = m.clean_rows(self._raw("", "x"))
        assert math.isnan(out.loc[0, "value"]) and out.loc[0, "value_status"] == "suppressed"

    def test_member_ids_come_from_coordinate_and_labels_are_trimmed(self):
        import clean_shelter_cost as m
        out = m.clean_rows(self._raw("120", ""))
        assert out.loc[0, "income_group_member_id"] == 13
        assert out.loc[0, "tenure_member_id"] == 2
        assert out.loc[0, "household_type"] == "Some category"
        assert out.loc[0, "value"] == 120


# ---------------------------------------------------------------------------
# Integration tests: real raw files -> temp staging dir -> validation
# ---------------------------------------------------------------------------

class TestPipelineOutputs:
    def test_raw_files_unchanged(self, cleaned_dir):
        assert cc.verify_raw_manifest() == []

    def test_all_five_outputs_written(self, cleaned_dir):
        for name in OUTPUT_FILES.values():
            assert (cleaned_dir / name).exists(), name

    def test_every_validation_check_passes(self, cleaned_dir):
        checks, _ = vcd.run_checks(staging_dir=cleaned_dir)
        failed = [f"{c.dataset}: {c.name} -- {c.detail}" for c in checks if not c.passed]
        assert not failed, "\n".join(failed)

    def test_census_population(self, cleaned_dir):
        df = pd.read_csv(cleaned_dir / OUTPUT_FILES["census_profile"])
        assert df.loc[df["characteristic_id"] == 1, "count_total"].item() == 422_630

    @pytest.mark.parametrize("tenure,expected", [("Owner", 13_005), ("Renter", 15_735)])
    def test_shelter_30_plus_counts(self, cleaned_dir, tenure, expected):
        df = pd.read_csv(cleaned_dir / OUTPUT_FILES["shelter_cost"])
        row = df[(df["income_group"] == "Total - Total income of household")
                 & (df["household_type"] == "Total - Household type including census family structure")
                 & (df["shelter_cost_ratio"] == "Spending 30% or more of income on shelter costs")
                 & (df["tenure"] == tenure)]
        assert row["value"].item() == expected

    def test_shelter_keeps_all_2389_rows(self, cleaned_dir):
        assert len(pd.read_csv(cleaned_dir / OUTPUT_FILES["shelter_cost"])) == 2389

    def test_starts_uses_windsor_not_ontario(self, cleaned_dir):
        df = pd.read_csv(cleaned_dir / OUTPUT_FILES["starts_completions"])
        assert set(df["source_geography_label"]) == {"Windsor"}
        total = df[(df["source_table"] == "Table A4_1") & (df["activity"] == "Starts") & (df["dwelling_type"] == "Total")]
        assert total["units"].item() == 83

    def test_rental_market_suppressed_rows_kept_as_nan(self, cleaned_dir):
        df = pd.read_csv(cleaned_dir / OUTPUT_FILES["rental_market"], keep_default_na=False, na_values=[""])
        sup = df[df["value_status"] == "suppressed"]
        assert len(sup) > 0 and sup["value"].isna().all() and (sup["raw_value"] == "**").all()

    def test_canada_has_windsor_and_only_cmas(self, cleaned_dir):
        df = pd.read_csv(cleaned_dir / OUTPUT_FILES["canada_rental_market"])
        assert "Windsor CMA" in set(df["centre"])
        assert not df["centre"].str.contains(r"10,000\+| CA$|Canada CMAs").any()
