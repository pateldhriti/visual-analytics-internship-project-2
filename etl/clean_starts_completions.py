"""
Clean Windsor CMA housing starts, completions and units under construction
from CMHC's Housing Information Monthly workbook, August 2026 edition.

Source: data/raw/cmhc/provincial-starts-completions-dwelling-type-08-26-en.xlsx
  Table A4_1  Starts and Completions by Dwelling Type, CMAs        August 2026
  Table A5_1  Under Construction by Dwelling Type, CMAs            August 2026
  Table A6_1  Cumulative Starts and Completions, CMAs              January to August 2026
Each lists CMAs in column A (rows 8-49) plus a "CMA Total" row; only the
'Windsor' row is extracted. Provincial and national rows are never used.

CMHC symbol '-' means "Nil or Zero" (workbook's Symbols sheet): stored as
units = 0 with value_flag 'nil_or_zero'. The dwelling-type components only sum
to the published Total when read this way (checked in validation).

Output: data/staging/windsor_starts_completions.csv -- one row per
(table x activity x dwelling type).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from cleaning_common import RAW_DIR, RAW_FILES, STAGING_DIR, assert_cells, normalize_text, parse_cmhc_starts_value, raw_text

OUTPUT_NAME = "windsor_starts_completions.csv"
SOURCE_GEOGRAPHY_LABEL = "Windsor"
DWELLING_TYPES = ["Singles", "Semis", "Row", "Apt. and Other", "Total"]

# sheet -> (title in row 2, period label in row 4, period_type, activities in column order)
TABLES = {
    "Table A4_1": ("Starts and Completions by Dwelling Type", "August 2026", "month", ["Starts", "Completions"]),
    "Table A5_1": ("Under Construction by Dwelling Type", "August 2026", "stock_at_month_end", ["Under Construction"]),
    "Table A6_1": ("Cumulative Starts and Completions by Dwelling Type", "January to August 2026", "year_to_date",
                   ["Starts", "Completions"]),
}
FIRST_DATA_ROW = 8


def find_geography_row(ws, label: str) -> int:
    matches = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if normalize_text(ws.cell(r, 1).value) == label]
    if len(matches) != 1:
        raise ValueError(f"{ws.title}: expected exactly one {label!r} row, found {len(matches)}")
    return matches[0]


def parse_table(ws, title: str, period: str, period_type: str, activities: list[str]) -> list[dict]:
    expected = {"A2": title, "A3": "Census Metropolitan Areas", "A4": period}
    for a, activity in enumerate(activities):
        for d, dwelling in enumerate(DWELLING_TYPES):
            col = 2 + a * len(DWELLING_TYPES) + d
            expected[f"{get_column_letter(col)}6"] = activity
            expected[f"{get_column_letter(col)}7"] = dwelling
    assert_cells(ws, expected)

    r = find_geography_row(ws, SOURCE_GEOGRAPHY_LABEL)
    rows = []
    for a, activity in enumerate(activities):
        for d, dwelling in enumerate(DWELLING_TYPES):
            col = 2 + a * len(DWELLING_TYPES) + d
            raw = ws.cell(r, col).value
            units, flag = parse_cmhc_starts_value(raw)
            rows.append({
                "geography": "Windsor CMA",
                "source_geography_label": SOURCE_GEOGRAPHY_LABEL,
                "source_table": ws.title,
                "activity": activity,
                "dwelling_type": dwelling,
                "period_label": period,
                "period_type": period_type,
                "units": units,
                "raw_value": raw_text(raw),
                "value_flag": flag,
                "source_cell": f"{ws.title}!{get_column_letter(col)}{r}",
            })
    return rows


def clean(raw_dir: Path = RAW_DIR, out_dir: Path = STAGING_DIR) -> dict:
    wb = openpyxl.load_workbook(raw_dir / RAW_FILES["starts_completions_xlsx"], data_only=True)
    rows = []
    for sheet, spec in TABLES.items():
        rows.extend(parse_table(wb[sheet], *spec))
    wb.close()

    df = pd.DataFrame(rows)
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates()
    out_dir.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_dir / OUTPUT_NAME, index=False, encoding="utf-8")
    return {
        "output": OUTPUT_NAME,
        "tables_read": list(TABLES),
        "raw_windsor_rows": len(TABLES),
        "cleaned_rows": len(df),
        "exact_duplicates_removed": duplicates,
        "value_flag_counts": df["value_flag"].value_counts().to_dict(),
    }


if __name__ == "__main__":
    for key, value in clean().items():
        print(f"{key}: {value}")
