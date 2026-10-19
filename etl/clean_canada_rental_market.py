"""
Clean CMA-level rental indicators from the CMHC Rental Market Survey 2025
national workbook, for later Windsor-vs-other-markets comparison.

Source: data/raw/cmhc/rmr-canada-2025-en.xlsx, Table 1.0 "Rental Market
Indicators - Privately Initiated Apartment Structures of Three Units and Over
- Provinces and Major Centres" (apartment-only scope).

Scope (agreed): every CMA row in Table 1.0, Windsor CMA included. Province
"10,000+" rows, census agglomerations (CA) and the "Canada CMAs" aggregate are
excluded. No comparator selection is made here -- all CMAs are kept.

Indicators, each for two periods:
  vacancy_rate (%)          Oct-24, Oct-25 (+ year-over-year trend marker)
  turnover_rate (%)         Oct-24, Oct-25 (+ trend marker)
  average_rent_2br ($)      Oct-24, Oct-25 (new and existing structures)
  rent_pct_change_2br (%)   Oct-23 to Oct-24, Oct-24 to Oct-25 (fixed sample)
'++' = change not statistically different from zero: value NaN,
value_status 'change_not_significant'. '**' = suppressed.

Output: data/staging/canada_rental_market_comparison.csv -- one row per
(centre x indicator x period).
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from cleaning_common import (
    RAW_DIR, RAW_FILES, STAGING_DIR, assert_cells, normalize_text, parse_cmhc_reliability,
    parse_cmhc_rms_value, parse_cmhc_trend, raw_text,
)

OUTPUT_NAME = "canada_rental_market_comparison.csv"
SHEET = "Table 1.0"
FIRST_DATA_ROW = 8
CMA_PATTERN = re.compile(r"\bCMA\b")  # matches "London CMA", "Ottawa-Gatineau CMA (Ont. part)"; not "Canada CMAs"

# metric -> (unit, [(value col, reliability col, period)], trend col or None)
METRICS = {
    "vacancy_rate": ("percent", [(2, 3, "Oct-24"), (4, 5, "Oct-25")], 6),
    "turnover_rate": ("percent", [(7, 8, "Oct-24"), (9, 10, "Oct-25")], 11),
    "average_rent_2br": ("dollars", [(12, 13, "Oct-24"), (14, 15, "Oct-25")], None),
    "rent_pct_change_2br": ("percent", [(16, 17, "Oct-23 to Oct-24"), (18, 19, "Oct-24 to Oct-25")], None),
}
HEADER = {
    "B6": "Vacancy Rates (%)",
    "G6": "Turnover Rates (%)",
    "L6": "Average Rent Two Bedroom ($) (New and Existing Structures)",
    "P6": "Percentage Change of Average Rent Two Bedroom From Fixed Sample (Existing Structures Only)",
    "A7": "Centre",
}


def data_rows(ws) -> list[int]:
    rows = []
    r = FIRST_DATA_ROW
    while normalize_text(ws.cell(r, 1).value) is not None:
        rows.append(r)
        r += 1
    return rows


def clean(raw_dir: Path = RAW_DIR, out_dir: Path = STAGING_DIR) -> dict:
    wb = openpyxl.load_workbook(raw_dir / RAW_FILES["canada_rms_xlsx"], data_only=True)
    ws = wb[SHEET]
    expected = dict(HEADER)
    for _, periods, _ in METRICS.values():
        for value_col, _, period in periods:
            expected[f"{get_column_letter(value_col)}7"] = period
    assert_cells(ws, expected)

    all_rows = data_rows(ws)
    cma_rows = [r for r in all_rows if CMA_PATTERN.search(normalize_text(ws.cell(r, 1).value))]
    records = []
    for r in cma_rows:
        centre = normalize_text(ws.cell(r, 1).value)
        for metric, (unit, periods, trend_col) in METRICS.items():
            for value_col, rel_col, period in periods:
                raw = ws.cell(r, value_col).value
                value, status = parse_cmhc_rms_value(raw)
                trend = None
                if trend_col and period == "Oct-25":
                    trend = parse_cmhc_trend(ws.cell(r, trend_col).value)
                records.append({
                    "source_table": SHEET,
                    "centre": centre,
                    "geography_level": "cma",
                    "metric": metric,
                    "unit": unit,
                    "period_label": period,
                    "value": value,
                    "raw_value": raw_text(raw),
                    "value_status": status,
                    "reliability_code": parse_cmhc_reliability(ws.cell(r, rel_col).value),
                    "yoy_trend": trend,
                    "source_cell": f"{SHEET}!{get_column_letter(value_col)}{r}",
                })
    wb.close()

    df = pd.DataFrame(records)
    if "Windsor CMA" not in set(df["centre"]):
        raise ValueError("Windsor CMA not found in Table 1.0")
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates()
    out_dir.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_dir / OUTPUT_NAME, index=False, encoding="utf-8")
    excluded = [normalize_text(ws.cell(r, 1).value) for r in all_rows if r not in cma_rows]
    return {
        "output": OUTPUT_NAME,
        "raw_centre_rows": len(all_rows),
        "cma_rows_kept": len(cma_rows),
        "rows_excluded_not_cma": excluded,
        "cleaned_rows": len(df),
        "exact_duplicates_removed": duplicates,
        "value_status_counts": df["value_status"].value_counts().to_dict(),
    }


if __name__ == "__main__":
    for key, value in clean().items():
        print(f"{key}: {value}")
