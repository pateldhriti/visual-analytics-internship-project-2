"""
Clean the CMHC Rental Market Survey 2025 workbook for Windsor CMA into one
tidy long table.

Source: data/raw/cmhc/rmr-windsor-2025-en.xlsx (29 sheets: contents + 28 tables)

Scope (agreed): the zone x bedroom-type tables for all three dwelling scopes,
plus the two Windsor-specific zone tables with their own layout:
  x.1.1 vacancy rate (%)        x.1.2 average rent ($)     x.1.3 rental universe (units)
  x.1.5 % change of average rent (fixed sample)            x.1.6 turnover rate (%)
  x.1.9 average rent ($) of vacant vs occupied units (October 2025)
  for x = 1 (apartment), 2 (row/townhouse), 3 (row and apartment combined)
  3.1.7 profile of new rental stock (structures completed July 2022-June 2025)
  3.1.8 vacancy rate and universe by renter household income quintile
Not extracted: the year-of-construction (1.2.x), structure-size (1.3.x) and
rent-quartile (1.4) tables.

Output: data/staging/windsor_rental_market.csv -- one row per published cell
(value + its reliability letter + trend/significance marker). Suppressed '**'
cells are kept as rows with value NaN and value_status 'suppressed'.

Every sheet's header cells are asserted before parsing (see
cleaning_common.assert_cells): a workbook with a different layout stops the
run rather than being silently misread.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from cleaning_common import (
    RAW_DIR, RAW_FILES, STAGING_DIR, assert_cells, normalize_text, parse_cmhc_reliability,
    parse_cmhc_rms_value, parse_cmhc_trend, raw_text,
)

OUTPUT_NAME = "windsor_rental_market.csv"
BEDROOM_TYPES = ["Studio", "1 Bedroom", "2 Bedroom", "3 Bedroom +", "Total"]
ZONE_LABELS = [
    "Zone 1 - Centre", "Zone 2 - East Inner", "Zone 3 - East Outer", "Zone 4 - West",
    "Windsor City (Zones 1-4)", "Zone 5 - Amherstburg", "Zone 6 - North Essex County",
    "Zone 7 - South Essex County", "Windsor CMA",
]
QUINTILES = ["N/A", "Q1", "Q2", "Q3", "Q4", "Q5", "Total"]
DWELLING_SCOPE = {"1": "apartment", "2": "row_townhouse", "3": "row_and_apartment"}

# table suffix -> (metric, unit, layout)
ZONE_TABLES = {
    "1.1": ("vacancy_rate", "percent", "rate_trend"),
    "1.2": ("average_rent", "dollars", "rent"),
    "1.3": ("rental_universe", "units", "universe"),
    "1.5": ("average_rent_pct_change", "percent", "pct_change"),
    "1.6": ("turnover_rate", "percent", "rate_trend"),
    "1.9": (None, "dollars", "vacant_occupied"),
}
# layout -> (group width, [(offset of value col, offset of reliability col or None, period label)])
LAYOUTS = {
    "rate_trend": (5, [(0, 1, "Oct-24"), (2, 3, "Oct-25")]),
    "rent": (4, [(0, 1, "Oct-24"), (2, 3, "Oct-25")]),
    "universe": (2, [(0, None, "Oct-24"), (1, None, "Oct-25")]),
    "pct_change": (4, [(0, 1, "Oct-23 to Oct-24"), (2, 3, "Oct-24 to Oct-25")]),
}
FIRST_ZONE_ROW = 8


def zone_level(zone: str) -> str:
    if zone == "Windsor CMA":
        return "cma_total"
    if zone.startswith("Windsor City"):
        return "city_subtotal"
    return "zone"


def cell_ref(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def record(ws, *, zone, bedroom, metric, unit, period, value_col, row, rel_col=None, **extra) -> dict:
    raw = ws.cell(row, value_col).value
    value, status = parse_cmhc_rms_value(raw)
    reliability = parse_cmhc_reliability(ws.cell(row, rel_col).value) if rel_col else None
    rec = {
        "source_table": ws.title,
        "zone": zone,
        "zone_level": zone_level(zone),
        "bedroom_type": bedroom,
        "metric": metric,
        "unit": unit,
        "period_label": period,
        "value": value,
        "raw_value": raw_text(raw),
        "value_status": status,
        "reliability_code": reliability,
        "yoy_trend": None,
        "vacant_rent_higher_flag": None,
        "income_quintile": None,
        "renter_income_range": None,
        "affordable_rent_range": None,
        "source_cell": f"{ws.title}!{cell_ref(value_col, row)}",
    }
    rec.update(extra)
    return rec


def check_zone_rows(ws, first_row: int) -> None:
    assert_cells(ws, {cell_ref(1, first_row + i): z for i, z in enumerate(ZONE_LABELS)})


def parse_zone_table(ws, metric: str, unit: str, layout: str) -> list[dict]:
    rows = []
    if layout == "vacant_occupied":
        expected = {"A7": "Zone"}
        for b, bedroom in enumerate(BEDROOM_TYPES):
            base = 2 + b * 5
            expected[cell_ref(base, 6)] = bedroom
            expected[cell_ref(base, 7)] = "Vacant Units"
            expected[cell_ref(base + 2, 7)] = "Occupied Units"
            expected[cell_ref(base + 4, 7)] = "Y/N"
        assert_cells(ws, expected)
        check_zone_rows(ws, FIRST_ZONE_ROW)
        for i, zone in enumerate(ZONE_LABELS):
            r = FIRST_ZONE_ROW + i
            for b, bedroom in enumerate(BEDROOM_TYPES):
                base = 2 + b * 5
                flag = parse_yn_flag(ws.cell(r, base + 4).value)
                for offset, name in ((0, "average_rent_vacant_units"), (2, "average_rent_occupied_units")):
                    rows.append(record(ws, zone=zone, bedroom=bedroom, metric=name, unit=unit, period="Oct-25",
                                       value_col=base + offset, rel_col=base + offset + 1, row=r,
                                       vacant_rent_higher_flag=flag))
        return rows

    width, periods = LAYOUTS[layout]
    expected = {"A7": "Zone"}
    for b, bedroom in enumerate(BEDROOM_TYPES):
        base = 2 + b * width
        expected[cell_ref(base, 6)] = bedroom
        for offset, _, period in periods:
            expected[cell_ref(base + offset, 7)] = period
    assert_cells(ws, expected)
    check_zone_rows(ws, FIRST_ZONE_ROW)

    for i, zone in enumerate(ZONE_LABELS):
        r = FIRST_ZONE_ROW + i
        for b, bedroom in enumerate(BEDROOM_TYPES):
            base = 2 + b * width
            for offset, rel_offset, period in periods:
                trend = None
                if layout == "rate_trend" and period == "Oct-25":
                    trend = parse_cmhc_trend(ws.cell(r, base + 4).value)
                rows.append(record(ws, zone=zone, bedroom=bedroom, metric=metric, unit=unit, period=period,
                                   value_col=base + offset,
                                   rel_col=None if rel_offset is None else base + rel_offset,
                                   row=r, yoy_trend=trend))
    return rows


def parse_yn_flag(raw) -> str | None:
    text = normalize_text(raw)
    if text is None:
        return None
    if text in ("Y", "N"):
        return text
    if text == "**":
        return "suppressed"
    raise ValueError(f"Unrecognised vacant-vs-occupied flag: {raw!r}")


def parse_new_stock(ws) -> list[dict]:
    """Table 3.1.7: per bedroom group [universe, vacancy value, vacancy
    reliability, average rent value, rent reliability]; zone rows 6-14."""
    expected = {"A5": "Zone"}
    for b, bedroom in enumerate(BEDROOM_TYPES):
        base = 2 + b * 5
        expected[cell_ref(base, 4)] = bedroom
        expected[cell_ref(base, 5)] = "Universe"
        expected[cell_ref(base + 1, 5)] = "Vacancy Rate"
        expected[cell_ref(base + 3, 5)] = "Average Rent"
    assert_cells(ws, expected)
    check_zone_rows(ws, 6)
    rows = []
    for i, zone in enumerate(ZONE_LABELS):
        r = 6 + i
        for b, bedroom in enumerate(BEDROOM_TYPES):
            base = 2 + b * 5
            common = dict(zone=zone, bedroom=bedroom, period="Oct-25", row=r)
            rows.append(record(ws, metric="new_stock_universe", unit="units", value_col=base, **common))
            rows.append(record(ws, metric="new_stock_vacancy_rate", unit="percent",
                               value_col=base + 1, rel_col=base + 2, **common))
            rows.append(record(ws, metric="new_stock_average_rent", unit="dollars",
                               value_col=base + 3, rel_col=base + 4, **common))
    return rows


def parse_income_quintile(ws) -> list[dict]:
    """Table 3.1.8: rows 7-69 = 9 zones x 7 quintile rows; per bedroom group
    (starting column E) [universe value, reliability, vacancy value, reliability]."""
    expected = {"A6": "Zone", "B6": "Income Quintile", "C6": "Renter household income range",
                "D6": "Monthly Rent Range Affordable"}
    for b, bedroom in enumerate(BEDROOM_TYPES):
        base = 5 + b * 4
        expected[cell_ref(base, 5)] = bedroom
        expected[cell_ref(base, 6)] = "Universe"
        expected[cell_ref(base + 2, 6)] = "Vacancy Rate"
    assert_cells(ws, expected)

    rows = []
    r = 7
    for zone in ZONE_LABELS:
        for quintile in QUINTILES:
            assert_cells(ws, {f"A{r}": zone, f"B{r}": quintile})
            extra = dict(income_quintile=quintile,
                         renter_income_range=normalize_text(ws.cell(r, 3).value),
                         affordable_rent_range=normalize_text(ws.cell(r, 4).value))
            for b, bedroom in enumerate(BEDROOM_TYPES):
                base = 5 + b * 4
                common = dict(zone=zone, bedroom=bedroom, period="Oct-25", row=r, **extra)
                rows.append(record(ws, metric="quintile_universe", unit="units",
                                   value_col=base, rel_col=base + 1, **common))
                rows.append(record(ws, metric="quintile_vacancy_rate", unit="percent",
                                   value_col=base + 2, rel_col=base + 3, **common))
            r += 1
    if normalize_text(ws.cell(r, 1).value) is not None:
        raise ValueError(f"{ws.title}: expected a blank row after the 63 data rows, found {ws.cell(r, 1).value!r}")
    return rows


def clean(raw_dir: Path = RAW_DIR, out_dir: Path = STAGING_DIR) -> dict:
    wb = openpyxl.load_workbook(raw_dir / RAW_FILES["windsor_rms_xlsx"], data_only=True, read_only=False)
    rows = []
    tables_read = []
    for scope_digit, scope in DWELLING_SCOPE.items():
        for suffix, (metric, unit, layout) in ZONE_TABLES.items():
            name = f"Table {scope_digit}.{suffix}"
            for rec in parse_zone_table(wb[name], metric, unit, layout):
                rec["dwelling_scope"] = scope
                rows.append(rec)
            tables_read.append(name)
    for name, parser in (("Table 3.1.7", parse_new_stock), ("Table 3.1.8", parse_income_quintile)):
        for rec in parser(wb[name]):
            rec["dwelling_scope"] = "row_and_apartment"
            rows.append(rec)
        tables_read.append(name)
    wb.close()

    df = pd.DataFrame(rows)
    df.insert(1, "dwelling_scope", df.pop("dwelling_scope"))
    df.insert(2, "geography", "Windsor CMA")
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates()

    out_dir.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_dir / OUTPUT_NAME, index=False, encoding="utf-8")
    return {
        "output": OUTPUT_NAME,
        "raw_tables_in_workbook": 28,
        "tables_read": tables_read,
        "cleaned_rows": len(df),
        "exact_duplicates_removed": duplicates,
        "value_status_counts": df["value_status"].value_counts().to_dict(),
        "rows_by_metric": df["metric"].value_counts().sort_index().to_dict(),
    }


if __name__ == "__main__":
    for key, value in clean().items():
        print(f"{key}: {value}")
