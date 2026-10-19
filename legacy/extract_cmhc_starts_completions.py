"""
Extract Ontario and Canada rows from CMHC's quarterly Starts, Completions,
and Under Construction by Dwelling Type workbook (Tables C1 and C2).

This CMHC table has no CMA-level breakdown -- it's published at the
Canada/provinces grain only (CMHC discontinued CMA-level completions
reporting for this series; see docs/conceptual_er_model.md). So unlike the
Windsor-specific shelter-cost and Census Profile extracts, this one keeps
Ontario (Windsor's province) and Canada (national benchmark) rather than
filtering to a single CMA -- the same "benchmark, not Windsor-specific"
role the Canada RMS workbook already plays for rental-market data.

Table C1 is the current-quarter snapshot ("Second Quarter 2026" etc.);
Table C2 is the year-to-date cumulative figure for the same tables'
columns. Both tables share the same (Activity, Province, dwelling-type
columns) layout, so both are read with the same function and tagged with
their own period_type/period_label.

Each table melts its 5 dwelling-type columns (Singles, Semis, Row,
Apartment and Other, Total) into rows, so one input row (one
activity x province) becomes 5 output rows (one per dwelling_type,
including the Total aggregate -- kept as an explicit row rather than
dropped, same convention as the other extracts in this project).

Run from the project virtual environment:
    .venv\\Scripts\\python.exe etl\\extract_cmhc_starts_completions.py
"""

from pathlib import Path

import openpyxl
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = PROJECT_ROOT / "data" / "raw" / "cmhc" / "quarterly-starts-completions-under-construction-06-26-en.xlsx"
OUTPUT_CSV = PROJECT_ROOT / "data" / "staging" / "ontario_canada_starts_completions.csv"

PROVINCE_FILTER = {"Ontario", "Canada"}
DWELLING_TYPE_COLUMNS = ["Singles", "Semis", "Row", "Apartment and Other", "Total"]

# sheet name -> period_type ('quarter' | 'ytd')
TABLES = {
    "Table C1": "quarter",
    "Table C2": "ytd",
}


def read_table(ws, period_type: str) -> pd.DataFrame:
    """Parse one CMHC Table Cx sheet: row 3 (1-indexed) is the period label
    ("Second Quarter 2026" / "January to June 2026"), row 5 is the column
    header, rows 6+ are (Activity, Province, dwelling-type counts...) data."""
    period_label = ws.cell(row=3, column=1).value
    header = [cell.value for cell in ws[5]]
    rows = [
        [cell.value for cell in row]
        for row in ws.iter_rows(min_row=6)
        if row[0].value is not None
    ]
    df = pd.DataFrame(rows, columns=header)
    df["period_type"] = period_type
    df["period_label"] = period_label
    return df


def melt_dwelling_types(df: pd.DataFrame) -> pd.DataFrame:
    """Melt the 5 dwelling-type count columns into (dwelling_type, units) rows."""
    id_cols = ["period_type", "period_label", "Activity", "Province"]
    long_df = df.melt(
        id_vars=id_cols,
        value_vars=DWELLING_TYPE_COLUMNS,
        var_name="dwelling_type",
        value_name="units",
    )
    return long_df.rename(columns={"Activity": "activity", "Province": "province"})


def extract(
    source_xlsx: Path = SOURCE_XLSX,
    output_csv: Path = OUTPUT_CSV,
    province_filter: set[str] = PROVINCE_FILTER,
) -> dict:
    if not source_xlsx.exists():
        raise FileNotFoundError(f"Source workbook not found: {source_xlsx}")

    wb = openpyxl.load_workbook(source_xlsx, data_only=True)

    frames = []
    total_source_rows = 0
    for sheet_name, period_type in TABLES.items():
        wide = read_table(wb[sheet_name], period_type)
        total_source_rows += len(wide)
        filtered = wide[wide["Province"].isin(province_filter)]
        frames.append(melt_dwelling_types(filtered))

    long_df = pd.concat(frames, ignore_index=True)
    long_df["units"] = pd.to_numeric(long_df["units"], errors="coerce")

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    long_df.to_csv(output_csv, index=False)

    return {
        "source_xlsx": str(source_xlsx),
        "output_csv": str(output_csv),
        "tables_read": list(TABLES),
        "province_filter": sorted(province_filter),
        "total_source_rows_both_tables": total_source_rows,
        "total_written_rows": len(long_df),
        "activities": sorted(long_df["activity"].unique()),
        "provinces_written": sorted(long_df["province"].unique()),
        "dwelling_types": sorted(long_df["dwelling_type"].unique()),
        "period_labels": sorted(long_df["period_label"].unique()),
    }


if __name__ == "__main__":
    stats = extract()
    print("Extraction complete.")
    for key, value in stats.items():
        print(f"  {key}: {value}")
